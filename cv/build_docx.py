from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Mm, Pt, RGBColor
from docx.oxml.ns import qn

OUT = Path(__file__).parent / "Shobhit-Gupta-CV.docx"


def set_font(style, name="Arial", size=10.4, bold=False, color="171717"):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_run(run, size=None, bold=None, color="171717"):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


doc = Document()
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Mm(12.5)
section.bottom_margin = Mm(12.5)
section.left_margin = Mm(15.5)
section.right_margin = Mm(15.5)

normal = doc.styles["Normal"]
set_font(normal, size=10.4)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(2.2)
normal.paragraph_format.line_spacing = 1.08

heading = doc.styles["Heading 1"]
set_font(heading, size=11.8, bold=True)
heading.paragraph_format.space_before = Pt(5.5)
heading.paragraph_format.space_after = Pt(1.2)
heading.paragraph_format.keep_with_next = True

if "CV Project Title" not in doc.styles:
    project_style = doc.styles.add_style("CV Project Title", WD_STYLE_TYPE.PARAGRAPH)
else:
    project_style = doc.styles["CV Project Title"]
set_font(project_style, size=10.6, bold=True)
project_style.paragraph_format.space_before = Pt(1)
project_style.paragraph_format.space_after = Pt(0.8)
project_style.paragraph_format.keep_with_next = True

# Name and contact block: ordinary paragraphs, not a table or layout grid.
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
set_run(p.add_run("Shobhit Gupta"), size=20, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0.5)
set_run(p.add_run("Full Stack Developer · Startup Builder · Problem Solver"), size=10.5)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0.2)
set_run(
    p.add_run(
        "shobhitguptacs25@bmsce.ac.in | https://github.com/shobhitguptacs | "
        "https://leetcode.com/u/Shobhit_Gupta_/"
    ),
    size=8.8,
)


def add_heading(text):
    return doc.add_paragraph(text, style="Heading 1")


def add_body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.widow_control = True
    return p


add_heading("Summary")
add_body(
    "Second-year CSE student at BMSCE and full-stack developer interested in product building and "
    "collaborative startup teams. Hands-on with React, Next.js, Node.js, Express, MongoDB, Python/Django, "
    "C++, JavaScript, data structures and algorithms, and basic machine learning with Python."
)

add_heading("Skills")
add_body(
    "Frontend: React, Next.js · Backend: Node.js, Express, Python, Django · Languages: C++, Python, "
    "JavaScript · Database: MongoDB · Problem solving: Data structures, algorithms · AI/ML: Basic "
    "machine learning with Python · Tools: Git, GitHub, Postman, Visual Studio Code"
)

add_heading("Projects")
doc.add_paragraph("PawCare — Stray Animal Assistance, Rescue and Adoption", style="CV Project Title")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1.5)
p.paragraph_format.keep_with_next = True
set_run(p.add_run("Technologies: "), size=9.6, bold=True)
set_run(
    p.add_run(
        "React, React Router, Express, MongoDB/Mongoose, JWT, bcrypt, Multer, Leaflet, OpenStreetMap"
    ),
    size=9.6,
)

bullet_style = doc.styles["List Bullet"]
set_font(bullet_style, size=10.2)
bullet_style.paragraph_format.left_indent = Mm(5)
bullet_style.paragraph_format.first_line_indent = Mm(-3.5)
bullet_style.paragraph_format.space_after = Pt(1.8)
bullet_style.paragraph_format.line_spacing = 1.05

bullets = [
    "Built PawCare end to end, a full-stack platform connecting people reporting stray or injured animals with volunteers, shelters, and adopters; added JWT/bcrypt authentication and role-protected views for users, volunteers, shelters, and administrators.",
    "Implemented photo- and map-based rescue reports with Emergency/Urgent/Normal priority, a volunteer Accept → In Progress → Resolved workflow, shelter-managed adoption listings, and lost/found notices.",
]
for text in bullets:
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.widow_control = True

add_heading("Education")
add_body("BMSCE — Computer Science and Engineering, second year.")

add_heading("Problem Solving")
add_body(
    "LeetCode: 30 problems solved (15 Easy, 13 Medium, 2 Hard); 21 active days; 7-day maximum streak. "
    "Profile checked 27 October 2026. Profile: https://leetcode.com/u/Shobhit_Gupta_/"
)

doc.core_properties.title = "Shobhit Gupta — Résumé"
doc.core_properties.subject = "One-page ATS-friendly résumé"
doc.core_properties.author = "Shobhit Gupta"
doc.save(OUT)
print(OUT)
